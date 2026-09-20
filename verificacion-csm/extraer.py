#!/usr/bin/env python3
"""Lee la planilla de Clínica Santa María y genera los datos de la aplicación.

Las fórmulas se interpretan automáticamente (no se transcriben a mano) para que
los coeficientes de dosis y los divisores de volumen sean exactamente los de la
planilla.
"""
import json, re, sys
import openpyxl

XLS = sys.argv[1] if len(sys.argv) > 1 else 'planilla_csm.xlsx'
wb = openpyxl.load_workbook(XLS, data_only=False)

def txt(v):
    if v is None:
        return ''
    if isinstance(v, str):
        return re.sub(r'\s+', ' ', v).strip()
    return str(v)

NUM = r'(\d+(?:\.\d+)?)'

def dosis(f, peso='B3', sc='D3'):
    """Devuelve ('kg', k) para k mg/Kg, o ('sc', k) para k por m2."""
    if not isinstance(f, str) or not f.startswith('='):
        return None
    s = f.replace(' ', '')
    for patron in (rf'^=\(?{NUM}\*{peso}\)?/1000$', rf'^=\(?{peso}\*{NUM}\)?/1000$'):
        m = re.match(patron, s)
        if m:
            return ('kg', float(m.group(1)))
    m = re.match(rf'^=\(?{sc}\*{NUM}\)?$', s)
    if m:
        return ('sc', float(m.group(1)))
    return None

def volumen(f, ref='E'):
    """Devuelve ('div', d) o ('mul', factor) a partir de la fórmula del volumen."""
    if not isinstance(f, str) or not f.startswith('='):
        return None
    s = f.replace(' ', '')
    m = re.match(rf'^=\(?{ref}\d+/{NUM}\)?$', s)
    if m:
        return ('div', float(m.group(1)))
    m = re.match(rf'^=\(?{ref}\d+\*{NUM}\)?/{NUM}$', s)
    if m:
        return ('mul', float(m.group(1)) / float(m.group(2)))
    m = re.match(rf'^=\(?{ref}\d+\*{NUM}/{NUM}\)?$', s)
    if m:
        return ('mul', float(m.group(1)) / float(m.group(2)))
    m = re.match(rf'^=\(?{ref}\d+\*{NUM}\)?$', s)
    if m:
        return ('mul', float(m.group(1)))
    return None

problemas = []

# ----------------------------------------------------------------- MED BOLO
ws = wb['MED BOLO']
bolos, actual = [], None
for r in range(9, 72):
    A, D, E = txt(ws[f'A{r}'].value), txt(ws[f'D{r}'].value), ws[f'E{r}'].value
    if A:
        actual = {
            'nombre': A, 'presentacion': txt(ws[f'B{r}'].value), 'via': txt(ws[f'C{r}'].value),
            'diluir': txt(ws[f'I{r}'].value), 'concentracion': txt(ws[f'J{r}'].value),
            'tiempo': txt(ws[f'K{r}'].value), 'nota': txt(ws[f'L{r}'].value), 'lineas': []
        }
        bolos.append(actual)
    if actual is None or not (D or E):
        continue
    d, v = dosis(E), volumen(ws[f'G{r}'].value)
    if E is not None and d is None:
        problemas.append(f'MED BOLO E{r}: {E!r}')
    if ws[f'G{r}'].value is not None and v is None:
        problemas.append(f'MED BOLO G{r}: {ws[f"G{r}"].value!r}')
    actual['lineas'].append({
        'celda': r,
        'etiqueta': txt(ws[f'C{r}'].value) if (not A and txt(ws[f'C{r}'].value)) else '',
        'dosisTxt': D, 'base': d[0] if d else None, 'k': d[1] if d else None,
        'unidad': txt(ws[f'F{r}'].value), 'vol': v[0] if v else None, 'f': v[1] if v else None,
        'xlDosis': txt(E), 'xlVol': txt(ws[f'G{r}'].value),
        'concLinea': txt(ws[f'J{r}'].value) if not A else '',
        'tiempoLinea': txt(ws[f'K{r}'].value) if not A else '',
    })

# ------------------------------------------------------------------ MED BIC
ws = wb['MED BIC']
bic, actual = [], None
for r in range(10, 27):
    A = txt(ws[f'A{r}'].value)
    H = ws[f'H{r}'].value
    if not (A or H):
        continue
    s = txt(H).replace(' ', '')
    m = re.match(rf'^=\(?B3\*E\d+\*(60\*24|24\*60|24)\)?/1000(/1000)?\*2$', s) or \
        re.match(rf'^=\(?B3\*(60\*24)\*E\d+\)?/1000(/1000)?\*2$', s)
    if not m:
        problemas.append(f'MED BIC H{r}: {H!r}')
        factor, div = None, None
    else:
        factor = 1440 if m.group(1) in ('60*24', '24*60') else 24
        div = 1000 if m.group(2) else 1
    vf = volumen(ws[f'J{r}'].value, 'H')
    if ws[f'J{r}'].value is not None and vf is None:
        problemas.append(f'MED BIC J{r}: {ws[f"J{r}"].value!r}')
    fila = {
        'celda': r,
        'dosisRec': txt(ws[f'C{r}'].value), 'preparar': txt(ws[f'D{r}'].value),
        'porCC': ws[f'E{r}'].value, 'unidad': txt(ws[f'F{r}'].value),
        'factor': factor, 'div': div, 'prepUnidad': txt(ws[f'I{r}'].value),
        'vialTipo': vf[0] if vf else None, 'vialF': vf[1] if vf else None,
        'xlPrep': txt(H), 'xlVial': txt(ws[f'J{r}'].value), 'nota': txt(ws[f'L{r}'].value),
    }
    if A:
        # En la planilla el nombre trae pegada la presentación y a veces una
        # nota, separadas por varios espacios: se parten para la lista.
        crudo = str(ws[f'A{r}'].value)
        partes = [x.strip() for x in re.split(r'\s{2,}', crudo) if x.strip()]
        actual = {'nombre': partes[0], 'detalle': ' · '.join(partes[1:]),
                  'concentracion': txt(ws[f'B{r}'].value), 'lineas': [fila]}
        bic.append(actual)
    elif actual:
        actual['lineas'].append(fila)

# --------------------------------------------------------------- MED ORALES
ws = wb['MED ORALES']
orales, actual = [], None
for r in range(9, 43):
    A, E, F = txt(ws[f'A{r}'].value), txt(ws[f'E{r}'].value), ws[f'F{r}'].value
    H = ws[f'H{r}'].value
    if not (A or E or H):
        continue
    d = dosis(F, peso='C3')
    if F is not None and d is None:
        problemas.append(f'MED ORALES F{r}: {F!r}')
    v = volumen(H, 'F')
    hDirecta = None
    if v is None and H is not None:
        dh = dosis(H, peso='C3')
        if dh:
            v, hDirecta = ('igual', 1.0), dh          # H se calcula desde el peso
        elif isinstance(H, (int, float)):
            v, hDirecta = ('fijo', float(H)), None
        elif isinstance(H, str) and not H.startswith('='):
            v, hDirecta = ('texto', H), None
        else:
            problemas.append(f'MED ORALES H{r}: {H!r}')
    fila = {
        'celda': r,
        'dosisTxt': E, 'base': d[0] if d else None, 'k': d[1] if d else None,
        'unidad': txt(ws[f'G{r}'].value), 'vol': v[0] if v else None,
        'f': v[1] if v else None, 'kDirecta': hDirecta[1] if hDirecta else None,
        'unidadVol': txt(ws[f'J{r}'].value), 'xlDosis': txt(F), 'xlVol': txt(H),
        'concLinea': txt(ws[f'C{r}'].value) if not A else '',
    }
    if A:
        actual = {'nombre': A, 'magistral': txt(ws[f'B{r}'].value) == 'SI',
                  'concentracion': txt(ws[f'C{r}'].value), 'via': txt(ws[f'D{r}'].value),
                  'nota': txt(ws[f'L{r}'].value), 'lineas': [fila]}
        orales.append(actual)
    elif actual:
        actual['lineas'].append(fila)

# Las filas que sólo llevan una aclaración de la dosis (p. ej. «30 mg/m2 /dia»)
# se pegan como detalle de la línea anterior en vez de quedar sueltas.
for grupo in (bolos, orales):
    for f in grupo:
        juntas = []
        for l in f['lineas']:
            if l.get('base') is None and l.get('vol') is None and l.get('dosisTxt') and juntas:
                juntas[-1]['detalle'] = l['dosisTxt']
            else:
                juntas.append(l)
        f['lineas'] = juntas

mapa = {'bolos': [[l['celda'] for l in f['lineas']] for f in bolos],
        'bic':   [[l['celda'] for l in f['lineas']] for f in bic],
        'orales':[[l['celda'] for l in f['lineas']] for f in orales]}
json.dump(mapa, open('celdas.json', 'w'))
for grupo in (bolos, bic, orales):
    for f in grupo:
        for l in f['lineas']:
            l.pop('celda', None)

def limpia(o):
    if isinstance(o, dict):
        return {k: limpia(v) for k, v in o.items() if v not in (None, '', [])}
    if isinstance(o, list):
        return [limpia(x) for x in o]
    return o

datos = {'BOLOS': limpia(bolos), 'BIC': limpia(bic), 'ORALES': limpia(orales)}
salida = ['\n'.join([
    '/* =====================================================================',
    '   Calculadora de medicamentos · Clínica Santa María',
    '   ARCHIVO GENERADO — no editar a mano.',
    '   Lo produce verificacion-csm/extraer.py leyendo la planilla original:',
    '   los coeficientes de dosis y los divisores de volumen salen de sus',
    '   fórmulas, no de una transcripción.',
    '   ===================================================================== */'])]
for nombre, valor in datos.items():
    salida.append(f'const {nombre}_CSM = ' + json.dumps(valor, ensure_ascii=False, indent=2) + ';')
open('../csm/js/datos-medicamentos.js', 'w').write('\n\n'.join(salida) + '\n')
print(f'csm/js/datos-medicamentos.js escrito', file=sys.stderr)
sys.stderr.write('\n'.join(['FÓRMULAS NO RECONOCIDAS:'] + problemas) + '\n' if problemas
                 else '\nTodas las fórmulas reconocidas.\n')
sys.stderr.write(f'bolos={len(bolos)} bic={len(bic)} orales={len(orales)}\n')
